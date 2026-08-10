#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verifica_publicacao_pncp.py  (v3 - janelas anuais + estado INDETERMINADO)
======================================================================
Verifica, para TODOS os municipios do Parana, se cada um JA PUBLICA
no PNCP (qualquer contratacao, qualquer modalidade) no periodo
01/01/2025 ate hoje.

Contexto legal: art. 176 da Lei 14.133/2021 concede aos municipios com
ate 20.000 habitantes prazo estendido para adesao ao PNCP.

CORRECOES IMPORTANTES DESTA VERSAO (v3):
  * A API /contratacoes/publicacao NAO aceita intervalos de data muito
    grandes (estoura em timeout/erro). Por isso a consulta e feita em
    JANELAS DE ATE 1 ANO (2025 inteiro; 2026 ate hoje), da mais recente
    para a mais antiga.
  * Falha de coleta (timeout/erro do PNCP) NAO e mais confundida com
    "nao publica". Agora ha TRES estados:
        SIM            -> achou pelo menos 1 publicacao
        NAO            -> consultou com sucesso e nao achou nada
        INDETERMINADO  -> a consulta falhou; nao da para afirmar
    (jamais marcar NAO quando na verdade a requisicao falhou.)

COMO FUNCIONA
  1) Baixa da API do IBGE a lista OFICIAL de TODOS os municipios do PR.
  2) Para cada municipio, testa as modalidades em cada janela anual.
     Ao achar o 1o registro -> marca SIM e passa ao proximo (agiliza).
  3) Gera 'publicacao_pncp_2026.xlsx' (NAO em vermelho, INDETERMINADO
     em amarelo) e imprime as listas no final.

Rode NA SUA MAQUINA (as APIs bloqueiam IPs de datacenter) OU gere o
.exe pelo GitHub Actions. Dependencias: requests, openpyxl
======================================================================
"""

import datetime as dt
import os
import re
import sys
import time
import unicodedata

try:
    import requests
except ImportError:
    sys.exit("Instale 'requests':  pip install requests")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    sys.exit("Instale 'openpyxl':  pip install openpyxl")

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ----------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------
BASE_PNCP = "https://pncp.gov.br/api/consulta"
IBGE_URL = "https://servicodados.ibge.gov.br/api/v1/localidades/estados/41/municipios"
DATA_INI = dt.date(2025, 1, 1)
DATA_FIM = dt.date.today()
ARQ_SAIDA = "publicacao_pncp_2026.xlsx"
TIMEOUT = 45
TENTATIVAS = 3
BACKOFF = 4
HEADERS = {"User-Agent": "Mozilla/5.0", "Accept": "application/json"}

# Modalidades mais usadas primeiro (para achar QUALQUER publicacao rapido).
#  6 Pregao Elet.| 8 Dispensa | 9 Inexig.| 4 Concorr.Elet.| 7 Pregao Pres.
#  5 Concorr.Pres.|12 Credenc.|13 Leilao Pres.|1 Leilao Elet.|2 Dialogo
#  3 Concurso |10 Manif.Interesse |11 Pre-qualif.
MODALIDADES = [6, 8, 9, 4, 7, 5, 12, 13, 1, 2, 3, 10, 11]


# ----------------------------------------------------------------------
def janelas_anuais(ini, fim):
    """Divide [ini, fim] em janelas de ate 1 ano, da MAIS RECENTE p/ a mais
    antiga (municipio ativo costuma ter publicacao recente -> acha rapido)."""
    js = []
    ano = fim.year
    while dt.date(ano, 1, 1) >= ini or ano == ini.year:
        a = max(ini, dt.date(ano, 1, 1))
        b = min(fim, dt.date(ano, 12, 31))
        if a <= b:
            js.append((a, b))
        if ano == ini.year:
            break
        ano -= 1
    return js   # ex.: [(2026-01-01, hoje), (2025-01-01, 2025-12-31)]


def norm(s):
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace("'", " ").replace("`", " ")
    s = re.sub(r"[^a-z0-9\s]", " ", s)
    return " ".join(s.split())


def get_municipios_pr():
    r = requests.get(IBGE_URL, timeout=TIMEOUT, headers=HEADERS)
    r.raise_for_status()
    muns = [(m["nome"], str(m["id"])) for m in r.json()]
    muns.sort(key=lambda x: norm(x[0]))
    return muns


# _get retorna:
#   ("ok", payload)   -> HTTP 200 com JSON
#   ("vazio", None)   -> HTTP 204 (sem conteudo, consulta valida)
#   ("erro", None)    -> falha (timeout / 5xx / 429 / outro) apos tentativas
def _get(endpoint, params):
    for tent in range(1, TENTATIVAS + 1):
        try:
            r = requests.get(f"{BASE_PNCP}{endpoint}", params=params,
                             timeout=TIMEOUT, headers=HEADERS)
        except requests.RequestException:
            if tent > 1:
                print(f"     >> PNCP instavel, aguarde... ({tent}/{TENTATIVAS})")
            time.sleep(BACKOFF * tent)
            continue
        if r.status_code == 204:
            return "vazio", None
        if r.status_code in (429, 500, 502, 503, 504):
            if tent > 1:
                print(f"     >> PNCP instavel, aguarde... ({tent}/{TENTATIVAS})")
            time.sleep(BACKOFF * tent)
            continue
        if r.status_code != 200:
            return "erro", None          # 400/422/etc: trata como indeterminado
        try:
            return "ok", r.json()
        except ValueError:
            return "erro", None
    return "erro", None


def verifica_municipio(ibge, janelas):
    """Retorna (status, modalidade, janela_str):
       status ∈ {'SIM','NAO','INDETERMINADO'}.
       - SIM: achou publicacao (para na 1a).
       - NAO: TODAS as consultas retornaram vazio com sucesso.
       - INDETERMINADO: nao achou nada, mas houve ao menos 1 falha de
         coleta (nao da para afirmar que nao publica)."""
    houve_erro = False
    for (a, b) in janelas:
        di, dfim = a.strftime("%Y%m%d"), b.strftime("%Y%m%d")
        for mod in MODALIDADES:
            params = {"dataInicial": di, "dataFinal": dfim,
                      "codigoModalidadeContratacao": mod, "uf": "PR",
                      "codigoMunicipioIbge": ibge,
                      "pagina": 1, "tamanhoPagina": 10}
            estado, payload = _get("/v1/contratacoes/publicacao", params)
            if estado == "erro":
                houve_erro = True
                continue
            if estado == "vazio":
                continue
            data = payload.get("data", payload.get("content", []))
            if data:
                return "SIM", mod, f"{a.year}"
            # 200 com lista vazia = consulta valida sem resultado -> continua
    return ("INDETERMINADO" if houve_erro else "NAO"), None, None


# ----------------------------------------------------------------------
def salva_excel(resultados):
    wb = Workbook()
    ws = wb.active
    ws.title = "Publicacao PNCP"
    headers = ["Municipio", "Codigo IBGE", "Publica no PNCP?",
               "Modalidade do 1o achado", "Ano do 1o achado",
               "Periodo verificado"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    verm = PatternFill("solid", fgColor="F8CBAD")   # NAO publica
    amar = PatternFill("solid", fgColor="FFF2CC")   # INDETERMINADO
    periodo = f"{DATA_INI:%d/%m/%Y} a {DATA_FIM:%d/%m/%Y}"
    for (nome, ibge, status, mod, ano) in resultados:
        ws.append([nome, ibge, status, mod or "-", ano or "-", periodo])
        if status == "NAO":
            fill = verm
        elif status == "INDETERMINADO":
            fill = amar
        else:
            fill = None
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill
    for i, wid in enumerate([28, 13, 18, 24, 16, 22], start=1):
        ws.column_dimensions[chr(64 + i)].width = wid
    ws.freeze_panes = "A2"
    caminho = os.path.join(BASE_DIR, ARQ_SAIDA)
    wb.save(caminho)
    return caminho


def main():
    janelas = janelas_anuais(DATA_INI, DATA_FIM)
    print("=" * 68)
    print(" VERIFICADOR DE PUBLICACAO NO PNCP (todos os municipios do PR)")
    print(f" Periodo: {DATA_INI:%d/%m/%Y} ate {DATA_FIM:%d/%m/%Y}")
    print(f" Janelas anuais consultadas: "
          f"{', '.join(f'{a.year}' for a, b in janelas)}")
    print("=" * 68)

    print("Baixando a lista oficial de municipios do PR (API do IBGE)...")
    try:
        municipios = get_municipios_pr()
    except Exception as e:
        sys.exit(f"[ERRO] Nao consegui a lista do IBGE: {e}\n"
                 "Rode este script na sua maquina (rede sem bloqueio).")
    print(f"  OK: {len(municipios)} municipios do PR.\n")

    resultados, nao, indet = [], [], []
    total = len(municipios)
    for i, (nome, ibge) in enumerate(municipios, start=1):
        status, mod, ano = verifica_municipio(ibge, janelas)
        etq = {"SIM": f"SIM (mod {mod}, {ano})",
               "NAO": "NAO PUBLICA",
               "INDETERMINADO": "INDETERMINADO (falha de coleta)"}[status]
        print(f"[{i:>3}/{total}] {nome:<30} {ibge}  -> {etq}")
        resultados.append((nome, ibge, status, mod, ano))
        if status == "NAO":
            nao.append((nome, ibge))
        elif status == "INDETERMINADO":
            indet.append((nome, ibge))
        time.sleep(0.15)

    caminho = salva_excel(resultados)

    print("\n" + "=" * 68)
    print(" RESULTADO")
    print("=" * 68)
    print(f"Total: {total}  |  Publicam: {total - len(nao) - len(indet)}  |  "
          f"NAO publicam: {len(nao)}  |  Indeterminados: {len(indet)}\n")
    print(f"Municipios que AINDA NAO PUBLICAM ({DATA_INI:%d/%m/%Y}"
          f" a {DATA_FIM:%d/%m/%Y}):")
    if nao:
        for nome, ibge in nao:
            print(f"   - {nome} ({ibge})")
    else:
        print("   (nenhum)")
    if indet:
        print(f"\nINDETERMINADOS (o PNCP falhou; REEXECUTE so estes depois):")
        for nome, ibge in indet:
            print(f"   - {nome} ({ibge})")
    print(f"\nPlanilha salva em:\n   {caminho}")

    if os.name == "nt":
        try:
            input("\nPressione ENTER para fechar...")
        except EOFError:
            pass


if __name__ == "__main__":
    main()
