#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
monitor_editais_pncp.py  (v8 - sem banco local; saida = planilha)
======================================================================
Monitor de editais do PNCP (UF=PR).

LOGICA (BUSCA DIRETA): voce cadastra na planilha a licitacao-alvo
  (MUNICIPIO, CODIGO IBGE, OBJETO, VALOR). A cada execucao o agente
  consulta o PNCP e sinaliza contratacao com:
       MESMO MUNICIPIO  E  ( OBJETO similar  OU  VALOR proximo )
  cabendo ao auditor avaliar retomada/republicacao x falso positivo.

SAIDA:
  - Se existir webhook do Teams (env/arquivo)  -> envia ao TEAMS.
  - Caso contrario -> gera PLANILHA de saida (.xlsx) com as licitacoes
    que deram match, para o auditor analisar. Colunas:
    Municipio | Objeto | Valor | Id da Contratacao (+ apoio).

v8: NAO cria mais banco local (editais_pncp.db). Cada execucao e
    independente e produz uma planilha com os candidatos atuais.

Planilha 'watchlist_editais.xlsx' (aba 'Editais'):
  A Ente(Nome Municipio) | B Codigo IBGE | C Objeto | D Valor |
  E Gerencia | F Observacoes

Robustez: consulta direta por municipio (codigoMunicipioIbge); timeout
60s; aviso "PNCP instavel, aguarde..."; retry/backoff em 429/5xx/timeout.

Conformidade API (Swagger pncp.gov.br/api/consulta/v3/api-docs):
  Base https://pncp.gov.br/api/consulta ; GET /v1/contratacoes/publicacao
  e /v1/contratacoes/atualizacao ; datas yyyyMMdd ;
  codigoModalidadeContratacao (OBRIGATORIO) ; uf ; codigoMunicipioIbge ;
  pagina (OBRIGATORIO) ; tamanhoPagina MAXIMO 50.
======================================================================
"""

import argparse
import datetime as dt
import os
import re
import sys
import time
import unicodedata
from difflib import SequenceMatcher

try:
    import requests
except ImportError:
    sys.exit("Falta a biblioteca 'requests' (embutida no .exe).")

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    sys.exit("Falta a biblioteca 'openpyxl' (embutida no .exe).")

try:
    from rapidfuzz import fuzz
    def obj_sim(a, b):
        return fuzz.token_set_ratio(a, b) / 100.0
except ImportError:
    def obj_sim(a, b):
        return SequenceMatcher(None, a, b).ratio()

try:
    import holidays
    _BR_HOLIDAYS = holidays.Brazil(subdiv="PR")
    def eh_dia_util(d):
        return d.weekday() < 5 and d not in _BR_HOLIDAYS
except ImportError:
    def eh_dia_util(d):
        return d.weekday() < 5


# ======================================================================
# LOCALIZACAO DOS ARQUIVOS (funciona como .py e como .exe do PyInstaller)
# ======================================================================
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def _carrega_env():
    for nome in (".env", "webhooks.env", "config.env"):
        caminho = os.path.join(BASE_DIR, nome)
        if os.path.exists(caminho):
            try:
                with open(caminho, "r", encoding="utf-8") as fh:
                    for linha in fh:
                        linha = linha.strip()
                        if not linha or linha.startswith("#") or "=" not in linha:
                            continue
                        k, v = linha.split("=", 1)
                        os.environ.setdefault(k.strip(), v.strip())
            except OSError:
                pass
            break


_carrega_env()


# ======================================================================
# CONFIGURACAO
# ======================================================================
BASE = "https://pncp.gov.br/api/consulta"
UF_ALVO = "PR"
JANELA_DIAS = 7
TAM_PAGINA = 50

TIMEOUT = 60
TENTATIVAS = 4
BACKOFF = 5
PAUSA_PAGINA = 0.5
HEADERS = {"User-Agent": "Mozilla/5.0", "Accept": "application/json"}

# Modalidades focadas em obras/engenharia (sem Pregao):
#   4 Concorrencia Elet. | 5 Concorrencia Pres. | 8 Dispensa | 9 Inexig.
MODALIDADES_OBRAS = [4, 5, 8, 9]

LIMIAR_OBJETO = 0.70
TOLERANCIA_VALOR = 0.20
SIM_MUNICIPIO = 0.90

XLSX = os.environ.get("WATCHLIST_XLSX",
                      os.path.join(BASE_DIR, "watchlist_editais.xlsx"))
ABA = os.environ.get("WATCHLIST_ABA", "Editais")

WEBHOOKS = {
    "INFRA":       os.environ.get("TEAMS_WEBHOOK_INFRA", "COLE_A_URL_DO_WEBHOOK_INFRA"),
    "EDIFICACOES": os.environ.get("TEAMS_WEBHOOK_EDIF",  "COLE_A_URL_DO_WEBHOOK_EDIFICACOES"),
}
GERENTES = {
    "INFRA":       "Lincoln Santos de Andrade",
    "EDIFICACOES": "Alexandre Cardoso Dal Ross",
}


def webhook_configurado():
    return any(u and not u.startswith("COLE_") for u in WEBHOOKS.values())


# ======================================================================
# TEXTO
# ======================================================================
STOPWORDS = {
    "de", "da", "do", "das", "dos", "e", "para", "com", "em", "no", "na",
    "a", "o", "as", "os", "um", "uma", "por", "que", "ao", "aos", "obra",
    "contratacao", "empresa", "especializada", "servicos", "execucao",
}


def normaliza(txt):
    if not txt:
        return ""
    txt = unicodedata.normalize("NFKD", str(txt))
    txt = "".join(c for c in txt if not unicodedata.combining(c)).lower()
    txt = re.sub(r"[^a-z0-9\s]", " ", txt)
    return " ".join(t for t in txt.split() if t not in STOPWORDS and len(t) > 2)


def normaliza_muni(txt):
    if not txt:
        return ""
    txt = unicodedata.normalize("NFKD", str(txt))
    txt = "".join(c for c in txt if not unicodedata.combining(c)).lower()
    txt = re.sub(r"[^a-z0-9\s]", " ", txt)
    tokens = [t for t in txt.split()
              if t not in ("municipio", "prefeitura", "de", "do", "da", "pr")]
    return " ".join(tokens).strip()


def so_digitos(s):
    return re.sub(r"\D", "", str(s or ""))


def norm_gerencia(g):
    g = normaliza(g).upper().replace(" ", "")
    if g.startswith("EDIF"):
        return "EDIFICACOES"
    if g.startswith("INFRA"):
        return "INFRA"
    return g


# ======================================================================
# WATCHLIST (Excel)
# ======================================================================
def _to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.-]", "", str(v)).replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def carrega_watchlist():
    if not os.path.exists(XLSX):
        print(f"[ERRO] Planilha nao encontrada: {XLSX}")
        print("       Coloque 'watchlist_editais.xlsx' na mesma pasta do programa.")
        return []
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    if ABA not in wb.sheetnames:
        print(f"[ERRO] Aba '{ABA}' nao existe. Abas: {wb.sheetnames}")
        return []
    ws = wb[ABA]
    watch = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 3:
            continue
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        ente   = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        ibge   = so_digitos(row[1])  if len(row) > 1 and row[1] is not None else ""
        objeto = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        valor  = _to_float(row[3] if len(row) > 3 else None)
        gerencia = norm_gerencia(row[4] if len(row) > 4 else "")
        if not objeto or gerencia not in ("INFRA", "EDIFICACOES"):
            continue
        ibge = ibge if len(ibge) == 7 else ""
        watch.append({
            "id": i, "ibge": ibge, "municipio": ente,
            "objeto": objeto, "valor": valor, "gerencia": gerencia,
        })
    wb.close()
    return watch


def cmd_list(args):
    watch = carrega_watchlist()
    if not watch:
        return
    print(f"\n{'LIN':<5}{'GER':<12}{'IBGE':<9}{'MUNICIPIO':<20}{'VALOR':>16}   OBJETO")
    print("-" * 96)
    for w in watch:
        v = f"R$ {w['valor']:,.2f}" if w["valor"] else "-"
        ibge = w["ibge"] or "(s/ IBGE)"
        print(f"{w['id']:<5}{w['gerencia']:<12}{ibge:<9}{(w['municipio'] or '')[:18]:<20}"
              f"{v:>16}   {w['objeto'][:34]}")
    print(f"\nTotal: {len(watch)} edital(is) vigiado(s).\n")


# ======================================================================
# COLETA NO PNCP
# ======================================================================
def _iso(d):
    return d.strftime("%Y%m%d")


def coleta(endpoint, data_ini, data_fim, modalidade, ibge=None):
    registros, pagina = [], 1
    while True:
        params = {"dataInicial": _iso(data_ini), "dataFinal": _iso(data_fim),
                  "codigoModalidadeContratacao": modalidade, "uf": UF_ALVO,
                  "pagina": pagina, "tamanhoPagina": TAM_PAGINA}
        if ibge:
            params["codigoMunicipioIbge"] = ibge
        r, sucesso = None, False
        for tent in range(1, TENTATIVAS + 1):
            try:
                r = requests.get(f"{BASE}{endpoint}", params=params,
                                 timeout=TIMEOUT, headers=HEADERS)
            except requests.RequestException as e:
                if tent == 1:
                    print(f"  ! rede em {endpoint} pag {pagina}: {e}")
                else:
                    print(f"  >> PNCP instavel, aguarde... (tentativa {tent}/{TENTATIVAS})")
                time.sleep(BACKOFF * tent)
                continue
            if r.status_code == 204:
                return registros
            if r.status_code in (429, 500, 502, 503, 504):
                if tent == 1:
                    print(f"  ! HTTP {r.status_code} em {endpoint} pag {pagina}")
                else:
                    print(f"  >> PNCP instavel, aguarde... (tentativa {tent}/{TENTATIVAS})")
                time.sleep(BACKOFF * tent)
                continue
            if r.status_code != 200:
                print(f"  ! HTTP {r.status_code} em {endpoint} pag {pagina}")
                return registros
            sucesso = True
            break
        if not sucesso:
            print(f"  ! desistindo da pag {pagina} de {endpoint} "
                  f"apos {TENTATIVAS} tentativas (PNCP instavel).")
            break
        payload = r.json()
        data = payload.get("data", payload.get("content", []))
        registros.extend(data)
        total_pag = payload.get("totalPaginas", payload.get("totalPages", 1))
        if pagina >= total_pag or not data:
            break
        pagina += 1
        time.sleep(PAUSA_PAGINA)
    return registros


def coleta_todos(watchlist, data_ini, data_fim):
    registros = []
    ibges = sorted({w["ibge"] for w in watchlist if w["ibge"]})
    tem_sem_ibge = any(not w["ibge"] for w in watchlist)
    if ibges:
        print(f"Consulta por municipio (IBGE): {', '.join(ibges)}")
        for cod in ibges:
            for mod in MODALIDADES_OBRAS:
                registros += coleta("/v1/contratacoes/atualizacao",
                                    data_ini, data_fim, mod, ibge=cod)
                registros += coleta("/v1/contratacoes/publicacao",
                                    data_ini, data_fim, mod, ibge=cod)
    if tem_sem_ibge:
        print("Ha linha(s) SEM IBGE -> varredura no estado inteiro (mais lenta).")
        for mod in MODALIDADES_OBRAS:
            print(f"  . estado, modalidade {mod} ...")
            registros += coleta("/v1/contratacoes/atualizacao",
                                data_ini, data_fim, mod)
            registros += coleta("/v1/contratacoes/publicacao",
                                data_ini, data_fim, mod)
    return registros


def extrai(item):
    uni = item.get("unidadeOrgao", {}) or {}
    org = item.get("orgaoEntidade", {}) or {}
    return {
        "id_pncp": item.get("numeroControlePNCP") or item.get("id"),
        "cnpj_orgao": org.get("cnpj"), "orgao": org.get("razaoSocial"),
        "municipio": uni.get("municipioNome"),
        "ibge": str(uni.get("codigoIbge") or uni.get("municipioIbge") or ""),
        "uf": uni.get("ufSigla"),
        "objeto": item.get("objetoCompra") or item.get("objeto") or "",
        "valor": item.get("valorTotalEstimado") or 0.0,
        "situacao": (item.get("situacaoCompraNome")
                     or str(item.get("situacaoCompraId") or "")),
        "data_public": item.get("dataPublicacaoPncp"),
        "data_atualiz": item.get("dataAtualizacao")
                        or item.get("dataAtualizacaoGlobal"),
    }


# ======================================================================
# MATCH DIRETO: municipio E (objeto OU valor)
# ======================================================================
def municipio_bate(watch, c):
    wi, ci = so_digitos(watch.get("ibge")), so_digitos(c.get("ibge"))
    if wi and ci:
        return wi == ci
    a = normaliza_muni(watch.get("municipio"))
    b = normaliza_muni(c.get("municipio"))
    if not a or not b:
        return False
    return a == b or obj_sim(a, b) >= SIM_MUNICIPIO


def valor_bate(watch, c):
    vw, vc = watch.get("valor"), c.get("valor")
    if not vw or not vc:
        return None
    return abs(vw - vc) / max(vw, vc) <= TOLERANCIA_VALOR


def avalia(watch, c):
    if not municipio_bate(watch, c):
        return False, 0.0, None, ""
    s_obj = obj_sim(normaliza(watch["objeto"]), normaliza(c["objeto"]))
    vflag = valor_bate(watch, c)
    obj_ok = s_obj >= LIMIAR_OBJETO
    val_ok = (vflag is True)
    if obj_ok and val_ok:
        return True, s_obj, vflag, "objeto+valor"
    if obj_ok:
        return True, s_obj, vflag, "objeto"
    if val_ok:
        return True, s_obj, vflag, "valor"
    return False, s_obj, vflag, ""


def detecta(registros, watchlist):
    """Sem banco: coleta os matches da execucao atual (dedup por
    (linha, id_pncp) para nao repetir quando o mesmo edital vem de
    /publicacao e /atualizacao)."""
    alertas, vistos = [], set()
    for item in registros:
        c = extrai(item)
        if not c["id_pncp"] or (c.get("uf") not in (UF_ALVO, None)):
            continue
        for w in watchlist:
            ok, s_obj, vflag, motivo = avalia(w, c)
            if not ok:
                continue
            chave = (w["id"], c["id_pncp"])
            if chave in vistos:
                continue
            vistos.add(chave)
            obs_val = ("valor n/d" if vflag is None
                       else (f"valor +-{int(TOLERANCIA_VALOR*100)}% OK" if vflag
                             else "valor fora"))
            det = (f"MATCH[{motivo}] linha {w['id']} | {c['orgao']} "
                   f"({c['municipio']}) | obj={s_obj:.0%} | "
                   f"R$ {c['valor']:,.2f} ({obs_val}) | "
                   f"situacao='{c['situacao']}' | {c['id_pncp']}")
            alertas.append(dict(w=w, c=c, s_obj=s_obj, vflag=vflag,
                                motivo=motivo, det=det))
    return alertas


# ======================================================================
# DIAGNOSTICO
# ======================================================================
def relatorio_watchlist(watchlist, registros):
    print("\n----- DIAGNOSTICO: situacao por edital vigiado -----")
    print("  [OK]=municipio E (objeto OU valor) | [~?]=municipio, mas nem "
          "objeto nem valor | [X]=municipio ausente na coleta")
    for w in watchlist:
        alvo = w["municipio"] or w["ibge"] or "?"
        achou_muni = False
        match_c, match_obj, match_motivo = None, 0.0, ""
        best_obj = 0.0
        for item in registros:
            c = extrai(item)
            if c.get("uf") not in (UF_ALVO, None):
                continue
            if not municipio_bate(w, c):
                continue
            achou_muni = True
            ok, s_obj, vflag, motivo = avalia(w, c)
            if s_obj > best_obj:
                best_obj = s_obj
            if ok and s_obj >= match_obj:
                match_c, match_obj, match_motivo = c, s_obj, motivo
        if match_c:
            print(f"  [OK] linha {w['id']} ({alvo}): disparo por {match_motivo} "
                  f"| obj={match_obj:.0%} | situacao='{match_c['situacao']}' "
                  f"| {match_c['id_pncp']}")
        elif achou_muni:
            print(f"  [~?] linha {w['id']} ({alvo}): ha edital(is) do municipio, "
                  f"mas nem objeto (melhor {best_obj:.0%}) nem valor bateram")
        else:
            print(f"  [X ] linha {w['id']} ({alvo}): nenhum edital deste "
                  f"municipio na coleta (verifique o codigo IBGE / janela)")
    print("----------------------------------------------------")


# ======================================================================
# PLANILHA DE SAIDA
# ======================================================================
def gera_planilha_saida(alertas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Licitacoes mapeadas"
    headers = ["Municipio", "Objeto", "Valor (R$)", "Id da Contratacao",
               "Alvo (linha)", "Disparo por", "Similaridade objeto",
               "Situacao", "Link PNCP"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    mot_map = {"objeto+valor": "objeto E valor", "objeto": "objeto similar",
               "valor": "valor proximo"}
    for a in alertas:
        c = a["c"]
        link = f"https://pncp.gov.br/app/editais/{c.get('id_pncp','')}"
        ws.append([c.get("municipio"), c.get("objeto"), c.get("valor", 0.0),
                   c.get("id_pncp"), a["w"]["id"],
                   mot_map.get(a["motivo"], a["motivo"]),
                   f"{a['s_obj']:.0%}", c.get("situacao"), link])
    widths = [20, 60, 16, 34, 11, 15, 18, 22, 42]
    for i, wid in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = wid
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = 'R$ #,##0.00'
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True,
                                                       vertical="top")
    ws.freeze_panes = "A2"
    nome = dt.datetime.now().strftime("saida_licitacoes_%Y%m%d_%H%M%S.xlsx")
    caminho = os.path.join(BASE_DIR, nome)
    wb.save(caminho)
    return caminho


# ======================================================================
# NOTIFICACAO NO TEAMS
# ======================================================================
def envia_teams(gerencia, c, s_obj, vflag, ref_linha, motivo="objeto/valor"):
    url = WEBHOOKS.get(gerencia, "")
    if not url or url.startswith("COLE_"):
        print(f"[AVISO] Webhook de {gerencia} nao configurado; alerta nao enviado.")
        return False
    link = f"https://pncp.gov.br/app/editais/{c.get('id_pncp','')}"
    obs_val = ("valor nao informado" if vflag is None
               else (f"dentro de +-{int(TOLERANCIA_VALOR*100)}%" if vflag
                     else "fora da faixa"))
    motivo_txt = {"objeto+valor": "objeto E valor",
                  "objeto": "objeto similar",
                  "valor": "valor proximo"}.get(motivo, motivo)
    facts = [
        {"title": "Possivel", "value": "retomada/republicacao (avaliar)"},
        {"title": "Disparado por", "value": motivo_txt},
        {"title": "Alvo (planilha)", "value": f"linha {ref_linha}"},
        {"title": "Ente", "value": f"{c.get('orgao','-')} ({c.get('municipio','-')})"},
        {"title": "Objeto (similaridade)", "value": f"{s_obj:.0%}"},
        {"title": "Valor", "value": f"R$ {c.get('valor',0):,.2f} ({obs_val})"},
        {"title": "Situacao", "value": c.get("situacao", "-")},
        {"title": "Gerencia", "value": f"{gerencia} - {GERENTES.get(gerencia,'')}"},
    ]
    card = {"type": "message", "attachments": [{
        "contentType": "application/vnd.microsoft.card.adaptive",
        "content": {"$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
            "type": "AdaptiveCard", "version": "1.4",
            "body": [
                {"type": "TextBlock", "size": "Large", "weight": "Bolder",
                 "color": "Warning",
                 "text": "Possivel retomada/republicacao de edital monitorado",
                 "wrap": True},
                {"type": "TextBlock", "text": c.get("objeto", "")[:280],
                 "wrap": True, "spacing": "Small"},
                {"type": "FactSet", "facts": facts}],
            "actions": [{"type": "Action.OpenUrl", "title": "Abrir no PNCP",
                         "url": link}]}}]}
    try:
        r = requests.post(url, json=card, timeout=30)
        ok = r.status_code in (200, 202)
        print(f"  -> Teams[{gerencia}] {'enviado' if ok else 'falhou'} (HTTP {r.status_code})")
        return ok
    except requests.RequestException as e:
        print(f"  ! erro Teams: {e}"); return False


def notifica(alertas):
    if not alertas:
        print(">> Nenhum edital do PNCP bateu com os alvos (municipio + objeto/valor).")
        return
    print(f"\n===== {len(alertas)} POSSIVEL(IS) MATCH(ES) - avaliar =====")
    for a in alertas:
        print(f"[MATCH] {a['det']}")

    if webhook_configurado():
        print("\nEnviando alertas ao Teams...")
        for a in alertas:
            envia_teams(a["w"]["gerencia"], a["c"], a["s_obj"],
                        a["vflag"], a["w"]["id"], a.get("motivo", "objeto/valor"))
    else:
        caminho = gera_planilha_saida(alertas)
        print("\n[SEM webhook] Planilha de saida gerada para analise "
              "(verdadeiro x falso positivo):")
        print(f"   {caminho}")


# ======================================================================
# EXECUCAO PRINCIPAL
# ======================================================================
def executar_monitoramento():
    hoje = dt.date.today()
    marca = "" if eh_dia_util(hoje) else "  (atencao: hoje nao e dia util BR/PR)"
    print("=" * 62)
    print(" MONITOR DE EDITAIS - PNCP (Parana)  [busca direta / por IBGE]")
    print(f" Execucao: {hoje:%d/%m/%Y}{marca}")
    print(f" Janela consultada: {(hoje - dt.timedelta(days=JANELA_DIAS)):%d/%m/%Y}"
          f" ate {hoje:%d/%m/%Y}")
    print(f" Criterio: mesmo MUNICIPIO E ( objeto >= {int(LIMIAR_OBJETO*100)}%"
          f"  OU  valor +-{int(TOLERANCIA_VALOR*100)}% )")
    print("=" * 62)

    watchlist = carrega_watchlist()
    if not watchlist:
        print("Watchlist vazia ou planilha ausente. Nada a fazer.")
        return
    print(f"Editais vigiados na planilha: {len(watchlist)}")

    ini = hoje - dt.timedelta(days=JANELA_DIAS)
    todos = coleta_todos(watchlist, ini, hoje)
    print(f"Registros coletados do PNCP: {len(todos)}")

    alertas = detecta(todos, watchlist)
    relatorio_watchlist(watchlist, todos)
    notifica(alertas)
    print("\nConcluido.")


def cmd_run(args):
    executar_monitoramento()


def cmd_testalert(args):
    g = norm_gerencia(args.gerencia or "INFRA")
    demo = {"id_pncp": "TESTE-0000", "orgao": "MUNICIPIO DE EXEMPLO",
            "municipio": "Curitiba", "valor": 1234567.89,
            "situacao": "Divulgada no PNCP",
            "objeto": "Teste de alerta do monitor de editais (PNCP)."}
    envia_teams(g, demo, 0.95, True, 0)


def _pausa_se_dois_cliques():
    if os.name == "nt":
        try:
            input("\nPressione ENTER para fechar...")
        except EOFError:
            pass


# ======================================================================
def main():
    if len(sys.argv) == 1:
        try:
            executar_monitoramento()
        except Exception as e:
            print(f"\n[ERRO] {e}")
        _pausa_se_dois_cliques()
        return

    p = argparse.ArgumentParser(description="Monitor de editais PNCP (PR) - busca direta por IBGE.")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list", help="listar a watchlist").set_defaults(func=cmd_list)
    rn = sub.add_parser("run", help="rodar varredura (ultimos 7 dias)")
    rn.set_defaults(func=cmd_run)
    t = sub.add_parser("testalert", help="testar webhook do Teams")
    t.add_argument("--gerencia", choices=["INFRA", "EDIFICACOES"])
    t.set_defaults(func=cmd_testalert)
    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
